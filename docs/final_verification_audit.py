import os
import re
import sys
import docx

docs_dir = r"D:\PROJECTS\ForensicsAnalyzer\docs"
files = [
    "1_Project_Building_Guide.docx",
    "2_Project_Student_Guide.docx",
    "3_Project_Student_Worksheet.docx",
    "4_Project_Student_Worksheet_Solutions.docx"
]

FORBIDDEN_PATTERNS = [
    "D:", "C:", ":\\", "\\", "127.0.0.1", "localhost", "5000", "http://",
    "SHA-256", "SHA256", "sha256"
]
HEX_32_REGEX = re.compile(r'[0-9a-fA-F]{32,}')

print("=== FINAL FORBIDDEN STRINGS AUDIT ON SAVED FILES ===\n")
total_violations = 0

for fname in files:
    fpath = os.path.join(docs_dir, fname)
    if not os.path.exists(fpath):
        print(f"ERROR: Missing file {fpath}")
        total_violations += 1
        continue

    doc = docx.Document(fpath)
    file_violations = [0]

    def check(text, loc):
        for pat in FORBIDDEN_PATTERNS:
            if pat in text:
                print(f"[{fname}] MATCH '{pat}' in {loc}: {text[:80]}")
                file_violations[0] += 1
        m = HEX_32_REGEX.search(text)
        if m:
            print(f"[{fname}] 32+ HEX MATCH in {loc}: {m.group(0)[:20]}...")
            file_violations[0] += 1

    # Check body paragraphs
    for i, p in enumerate(doc.paragraphs):
        check(p.text, f"Para {i+1}")

    # Check tables
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, p in enumerate(cell.paragraphs):
                    check(p.text, f"Tbl {t_idx+1} R{r_idx+1}C{c_idx+1} P{p_idx+1}")

    # Check headers & footers
    for s_idx, sec in enumerate(doc.sections):
        for p in sec.header.paragraphs:
            check(p.text, f"Sec {s_idx+1} Header")
        for table in sec.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        check(p.text, f"Sec {s_idx+1} Header Tbl")
        for p in sec.footer.paragraphs:
            check(p.text, f"Sec {s_idx+1} Footer")
        for table in sec.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        check(p.text, f"Sec {s_idx+1} Footer Tbl")

    sz = os.path.getsize(fpath)
    if file_violations[0] == 0:
        print(f"[PASS] {fname:<42}: 0 forbidden terms ({sz:,} bytes)")
    else:
        print(f"[FAIL] {fname:<42}: {file_violations[0]} violations found!")
        total_violations += file_violations[0]

print("\n" + "="*50)
if total_violations == 0:
    print("ALL 4 ON-DISK DOCX FILES ARE 100% CLEAN AND VERIFIED!")
else:
    print(f"AUDIT FAILED: {total_violations} TOTAL VIOLATIONS!")
    sys.exit(1)
